from dataclasses import dataclass
from io import BytesIO
from pathlib import Path

from docx import Document as DocxDocument
from pypdf import PdfReader

from app.models.document import Document
from app.services.object_storage import object_storage


SUPPORTED_EXTENSIONS = {
    ".txt",
    ".md",
    ".pdf",
    ".docx",
}


@dataclass
class LoadedDocument:
    document_id: str
    filename: str
    content_type: str | None
    text: str


def _read_object_content(document: Document) -> bytes:
    response = object_storage.client.get_object(
        bucket_name=object_storage.bucket_name,
        object_name=document.object_name,
    )

    try:
        return response.read()
    finally:
        response.close()
        response.release_conn()


def _extract_text_file(content: bytes) -> str:
    try:
        return content.decode("utf-8")
    except UnicodeDecodeError:
        return content.decode("latin-1")


def _extract_pdf(content: bytes) -> str:
    reader = PdfReader(BytesIO(content))

    pages = [
        page.extract_text().strip()
        for page in reader.pages
        if page.extract_text()
    ]

    return "\n\n".join(pages)


def _extract_docx(content: bytes) -> str:
    document = DocxDocument(BytesIO(content))

    paragraphs = [
        paragraph.text.strip()
        for paragraph in document.paragraphs
        if paragraph.text.strip()
    ]

    return "\n\n".join(paragraphs)


def load_document(document: Document) -> LoadedDocument:
    extension = Path(document.filename).suffix.lower()

    if extension not in SUPPORTED_EXTENSIONS:
        raise ValueError(
            f"Formato no soportado para extracción: {extension or 'sin extensión'}"
        )

    content = _read_object_content(document)

    if extension in {".txt", ".md"}:
        text = _extract_text_file(content)
    elif extension == ".pdf":
        text = _extract_pdf(content)
    elif extension == ".docx":
        text = _extract_docx(content)
    else:
        raise ValueError(
            f"No existe extractor para el formato: {extension}"
        )

    text = text.strip()

    if not text:
        raise ValueError(
            f"No se pudo extraer texto del documento: {document.filename}"
        )

    return LoadedDocument(
        document_id=str(document.id),
        filename=document.filename,
        content_type=document.content_type,
        text=text,
    )


def load_text_document(document: Document) -> LoadedDocument:
    """
    Alias temporal para mantener compatibilidad con el pipeline existente.
    """

    return load_document(document)
